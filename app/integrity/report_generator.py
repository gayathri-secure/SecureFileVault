
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ReportGenerationError(Exception):
    """Raised when an Excel integrity report cannot be generated."""


ENCRYPTION_HEADERS = [
    "Original File",
    "Encrypted File",
    "Original SHA-256",
    "Encrypted SHA-256",
    "Original Size (Bytes)",
    "Encrypted Size (Bytes)",
    "Algorithm",
    "Status",
]

DECRYPTION_HEADERS = [
    "Encrypted File",
    "Restored File",
    "Original SHA-256",
    "Restored SHA-256",
    "Restored Size (Bytes)",
    "Format Version",
    "Algorithm",
    "Integrity Status",
    "Decryption Status",
]


def _validate_hash(value, field_name, allow_none=False):
    if value is None and allow_none:
        return

    if (
        not isinstance(value, str)
        or len(value) != 64
        or any(
            char not in "0123456789abcdefABCDEF"
            for char in value
        )
    ):
        raise ReportGenerationError(
            f"Invalid SHA-256 value for {field_name}."
        )


def _style_report(worksheet, header_row, headers, widths):
    for cell in worksheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1D4ED8",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(headers))}"
        f"{worksheet.max_row}"
    )
    worksheet.row_dimensions[header_row].height = 36

    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[
            get_column_letter(index)
        ].width = width


def _save_workbook(workbook, destination):
    if destination.suffix.lower() != ".xlsx":
        raise ReportGenerationError(
            "The report filename must end with .xlsx."
        )

    if not destination.parent.is_dir():
        raise ReportGenerationError(
            "The report destination folder does not exist."
        )

    try:
        with destination.open("xb") as output_file:
            workbook.save(output_file)
    except FileExistsError as exc:
        raise ReportGenerationError(
            "The report already exists. Choose a new filename."
        ) from exc
    except (OSError, ValueError) as exc:
        raise ReportGenerationError(
            f"Could not save the Excel report: {exc}"
        ) from exc


def generate_encryption_report(
    report_path,
    records,
    algorithm="AES-256-GCM",
):
    """Generate an Excel report for encrypted files."""

    destination = Path(report_path)

    if not isinstance(records, (list, tuple)):
        raise ReportGenerationError(
            "Report records must be a list or tuple."
        )

    if not isinstance(algorithm, str) or not algorithm.strip():
        raise ReportGenerationError(
            "An encryption algorithm is required."
        )

    required_fields = {
        "original_path",
        "encrypted_path",
        "original_size_bytes",
        "encrypted_size_bytes",
        "original_sha256",
        "encrypted_sha256",
    }

    validated_records = []

    for index, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise ReportGenerationError(
                f"Record {index} must be a dictionary."
            )

        missing = required_fields - record.keys()
        if missing:
            raise ReportGenerationError(
                f"Record {index} is missing fields: "
                + ", ".join(sorted(missing))
            )

        for field in (
            "original_size_bytes",
            "encrypted_size_bytes",
        ):
            value = record[field]
            if (
                not isinstance(value, int)
                or isinstance(value, bool)
                or value < 0
            ):
                raise ReportGenerationError(
                    f"Record {index} has an invalid {field}."
                )

        _validate_hash(record["original_sha256"], "original_sha256")
        _validate_hash(record["encrypted_sha256"], "encrypted_sha256")

        validated_records.append(record)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Encryption Report"

    worksheet.append([
        "SecureFileVault - Encryption Integrity Report"
    ])
    worksheet.append([
        "Generated (UTC)",
        datetime.now(timezone.utc).isoformat(timespec="seconds"),
    ])
    worksheet.append(["Algorithm", algorithm])
    worksheet.append([])
    worksheet.append(ENCRYPTION_HEADERS)

    header_row = 5

    for record in validated_records:
        worksheet.append([
            str(record["original_path"]),
            str(record["encrypted_path"]),
            record["original_sha256"].lower(),
            record["encrypted_sha256"].lower(),
            record["original_size_bytes"],
            record["encrypted_size_bytes"],
            algorithm,
            "Encrypted",
        ])

    _style_report(
        worksheet,
        header_row,
        ENCRYPTION_HEADERS,
        [35, 45, 68, 68, 22, 24, 25, 16],
    )

    _save_workbook(workbook, destination)

    return {
        "report_path": str(destination.resolve()),
        "record_count": len(validated_records),
        "algorithm": algorithm,
    }


def generate_decryption_report(
    report_path,
    records,
    algorithm="AES-256-GCM",
):
    """
    Generate an Excel report for decrypted files.

    Each record should contain:
        encrypted_path
        restored_path
        original_sha256
        restored_sha256
        size_bytes
        format_version
        integrity_verified
        integrity_status
    """

    destination = Path(report_path)

    if not isinstance(records, (list, tuple)):
        raise ReportGenerationError(
            "Report records must be a list or tuple."
        )

    if not isinstance(algorithm, str) or not algorithm.strip():
        raise ReportGenerationError(
            "An encryption algorithm is required."
        )

    required_fields = {
        "encrypted_path",
        "restored_path",
        "original_sha256",
        "restored_sha256",
        "size_bytes",
        "format_version",
        "integrity_verified",
        "integrity_status",
    }

    validated_records = []

    for index, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise ReportGenerationError(
                f"Record {index} must be a dictionary."
            )

        missing = required_fields - record.keys()
        if missing:
            raise ReportGenerationError(
                f"Record {index} is missing fields: "
                + ", ".join(sorted(missing))
            )

        _validate_hash(
            record["original_sha256"],
            "original_sha256",
            allow_none=True,
        )
        _validate_hash(
            record["restored_sha256"],
            "restored_sha256",
        )

        size = record["size_bytes"]
        if (
            not isinstance(size, int)
            or isinstance(size, bool)
            or size < 0
        ):
            raise ReportGenerationError(
                f"Record {index} has an invalid size_bytes."
            )

        if record["format_version"] not in (1, 2):
            raise ReportGenerationError(
                f"Record {index} has an invalid format_version."
            )

        if not isinstance(record["integrity_verified"], bool):
            raise ReportGenerationError(
                f"Record {index} has an invalid integrity_verified value."
            )

        if not isinstance(record["integrity_status"], str):
            raise ReportGenerationError(
                f"Record {index} has an invalid integrity_status."
            )

        if record["format_version"] == 2:
            if record["integrity_verified"] is not True:
                raise ReportGenerationError(
                    f"Record {index}: SFV2 integrity was not verified."
                )

            if record["original_sha256"] is None:
                raise ReportGenerationError(
                    f"Record {index}: SFV2 original hash is missing."
                )

            if (
                record["original_sha256"].lower()
                != record["restored_sha256"].lower()
            ):
                raise ReportGenerationError(
                    f"Record {index}: original and restored hashes differ."
                )

        validated_records.append(record)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Decryption Report"

    worksheet.append([
        "SecureFileVault - Decryption Integrity Report"
    ])
    worksheet.append([
        "Generated (UTC)",
        datetime.now(timezone.utc).isoformat(timespec="seconds"),
    ])
    worksheet.append(["Algorithm", algorithm])
    worksheet.append([])
    worksheet.append(DECRYPTION_HEADERS)

    header_row = 5

    for record in validated_records:
        worksheet.append([
            str(record["encrypted_path"]),
            str(record["restored_path"]),
            (
                record["original_sha256"].lower()
                if record["original_sha256"]
                else "NOT AVAILABLE"
            ),
            record["restored_sha256"].lower(),
            record["size_bytes"],
            f"SFV{record['format_version']}",
            algorithm,
            record["integrity_status"],
            "Decrypted",
        ])

    _style_report(
        worksheet,
        header_row,
        DECRYPTION_HEADERS,
        [45, 45, 68, 68, 22, 16, 22, 28, 20],
    )

    _save_workbook(workbook, destination)

    return {
        "report_path": str(destination.resolve()),
        "record_count": len(validated_records),
        "algorithm": algorithm,
    }
